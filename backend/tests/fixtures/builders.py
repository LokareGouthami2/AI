"""Build fixture documents on the fly (no binary fixtures committed)."""

from __future__ import annotations

import io

import pymupdf as fitz  # PyMuPDF
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from backend.tests.fixtures.sample_content import ML_CHAPTER

_STYLES = {
    "TITLE": ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=22, leading=28, spaceAfter=14),
    "HEADING": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=16, leading=20, spaceBefore=10, spaceAfter=6),
    "SUBHEADING": ParagraphStyle("s", fontName="Helvetica-Bold", fontSize=13, leading=16, spaceBefore=6, spaceAfter=4),
    "FORMULA": ParagraphStyle("f", fontName="Courier", fontSize=11, leading=14, leftIndent=36, spaceAfter=6),
    "REFERENCE": ParagraphStyle("r", fontName="Helvetica", fontSize=9, leading=12, spaceAfter=3),
}
_BODY = ParagraphStyle("b", fontName="Helvetica", fontSize=11, leading=15, spaceAfter=8)


def _escape(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(content: list[tuple[str, str]] = ML_CHAPTER, repeat: int = 1) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="Sample")
    story = []
    for _ in range(repeat):
        for label, t in content:
            style = _STYLES.get(label, _BODY)
            # Helvetica lacks Σ and ŷ; keep formula glyphs ASCII-safe for the fixture.
            if label == "FORMULA":
                t = t.replace("Σ", "sum").replace("ŷ", "yhat").replace("−", "-").replace("²", "^2")
            story.append(Paragraph(_escape(t), style))
        story.append(Spacer(1, 12))
    doc.build(story)
    return buf.getvalue()


def build_scanned_pdf(dpi: int = 150) -> bytes:
    """Rasterise the text PDF so pages have no text layer (simulates a scan)."""
    src = fitz.open(stream=build_pdf(), filetype="pdf")
    out = fitz.open()
    for page in src:
        pix = page.get_pixmap(dpi=dpi)
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, stream=pix.tobytes("png"))
    data = out.tobytes()
    out.close()
    src.close()
    return data


def build_docx(content: list[tuple[str, str]] = ML_CHAPTER) -> bytes:
    import docx

    d = docx.Document()
    for label, t in content:
        if label == "TITLE":
            d.add_heading(t, level=0)
        elif label == "HEADING":
            d.add_heading(t, level=1)
        elif label == "SUBHEADING":
            d.add_heading(t, level=2)
        else:
            d.add_paragraph(t)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_txt(content: list[tuple[str, str]] = ML_CHAPTER) -> bytes:
    return "\n\n".join(t for _, t in content).encode("utf-8")
