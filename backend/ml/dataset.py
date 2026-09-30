"""Synthetic-but-realistic training data for the section classifier.

Pipeline:  corpus frames → labelled document spec → rendered as PDF / DOCX / TXT
with a randomised typographic style → *our real extraction + NLP pipeline* →
segments aligned back to their source labels.

Because labels are known by construction there is no annotation noise, and
because the features come from the production extractor they match inference
exactly. Splits are made by SUBJECT (no subject appears in two splits), which
prevents vocabulary leakage. Limitations are documented in docs/ml-pipeline.md.
"""

from __future__ import annotations

import io
import json
import random
import re
from dataclasses import dataclass, field
from pathlib import Path

from backend.documents.extract import extract
from backend.ml import corpus
from backend.nlp.preprocess import preprocess

SYS_FONTS = Path("/usr/share/fonts/truetype")


@dataclass
class Style:
    family: str
    body_size: float
    title_size: float
    heading_size: float
    sub_size: float
    heading_bold: bool
    sub_bold: bool
    heading_numbering: str  # "none" | "arabic" | "chapter"
    question_prefix: bool
    ref_style: str  # "apa" | "ieee"
    fmt: str  # pdf | docx | txt


@dataclass
class DocSpec:
    doc_id: str
    subject: str
    style: Style
    blocks: list[tuple[str, str]] = field(default_factory=list)


def _cap(s: str) -> str:
    return s[:1].upper() + s[1:] if s else s


def _bare(term: str) -> str:
    return re.sub(r"^(the|an|a)\s+", "", term)


class Generator:
    def __init__(self, seed: int):
        self.r = random.Random(seed)

    def pick(self, seq):
        return self.r.choice(seq)

    def _slots(self, subj: dict, subject_name: str) -> dict:
        (term, definition), (term2, _) = self.r.sample(subj["terms"], 2)
        process = self.pick(subj["processes"])
        # "regulates the flow" -> base form "regulate the flow" (crude but fine for frames)
        first, _, rest = process.partition(" ")
        base = (first[:-1] if first.endswith("s") and not first.endswith("ss") else first) + (" " + rest if rest else "")
        if first.endswith("ies"):
            base = first[:-3] + "y" + (" " + rest if rest else "")
        elif first.endswith(("shes", "ches", "sses", "xes")):
            base = first[:-2] + (" " + rest if rest else "")
        return {
            "term": term,
            "Term": _cap(term),
            "term_bare": _bare(term),
            "Term_bare": _cap(_bare(term)),
            "term2": term2,
            "term2_bare": _bare(term2),
            "definition": definition,
            "process": process,
            "process_base": base,
            "subject": subject_name.replace("_", " "),
        }

    def _fill(self, frame: str, slots: dict, **extra) -> str:
        return frame.format(**{**slots, **extra})

    def paragraph(self, subj, name) -> str:
        n = self.r.choice([1, 2, 2, 3])
        return " ".join(self._fill(self.pick(corpus.PARAGRAPH_FRAMES), self._slots(subj, name)) for _ in range(n))

    def definition(self, subj, name) -> str:
        return self._fill(self.pick(corpus.DEFINITION_FRAMES), self._slots(subj, name))

    def example(self, subj, name) -> str:
        s = self._slots(subj, name)
        body = self._fill(self.pick(corpus.EXAMPLE_BODIES), s)
        lc = body[:1].lower() + body[1:]
        return self._fill(self.pick(corpus.EXAMPLE_FRAMES), s, example_text=body, example_text_lc=lc)

    def procedure(self) -> str:
        a, b, c = self.r.sample(corpus.PROCEDURE_STEPS, 3)
        lc = lambda t: t[:1].lower() + t[1:]  # noqa: E731
        return self.pick(corpus.PROCEDURE_FRAMES).format(a=a, b=b, c=c, a_lc=lc(a), b_lc=lc(b), c_lc=lc(c))

    def important(self, subj, name) -> str:
        return self._fill(self.pick(corpus.IMPORTANT_FRAMES), self._slots(subj, name))

    def question(self, subj, name, n: int, prefix: bool) -> str:
        frames = corpus.QUESTION_FRAMES if prefix else corpus.QUESTION_FRAMES[4:]
        return self._fill(self.pick(frames), self._slots(subj, name), n=n)

    def answer(self, subj, name, n: int) -> str:
        return self._fill(self.pick(corpus.ANSWER_FRAMES), self._slots(subj, name), n=n)

    def conclusion(self, subj, name) -> str:
        return self._fill(self.pick(corpus.CONCLUSION_FRAMES), self._slots(subj, name))

    def reference(self, subj, idx: int, style: str) -> str:
        author, title, pub, year = subj["people"][idx % len(subj["people"])]
        if style == "ieee":
            return f"[{idx + 1}] {author}, {title}. {pub}, {year}."
        return f"{author} ({year}). {title}. {pub}."

    def style(self, fmt: str) -> Style:
        body = self.pick([10.0, 10.5, 11.0, 11.0, 12.0])
        return Style(
            family=self.pick(["sans", "serif", "sans", "dejavu"]),
            body_size=body,
            title_size=body + self.pick([8, 10, 12, 14]),
            heading_size=body + self.pick([3, 4, 5, 6]),
            sub_size=body + self.pick([1, 1.5, 2]),
            heading_bold=self.r.random() < 0.9,
            sub_bold=self.r.random() < 0.75,
            heading_numbering=self.pick(["none", "arabic", "arabic", "chapter"]),
            question_prefix=self.r.random() < 0.7,
            ref_style=self.pick(["apa", "ieee"]),
            fmt=fmt,
        )

    def document(self, doc_id: str, subject_name: str, fmt: str) -> DocSpec:
        subj = corpus.SUBJECTS[subject_name]
        st = self.style(fmt)
        spec = DocSpec(doc_id=doc_id, subject=subject_name, style=st)
        b = spec.blocks
        b.append(("TITLE", self.pick(subj["titles"])))
        topics = self.r.sample(subj["topics"], self.r.randint(2, 4))
        if self.r.random() < 0.5:
            topics.insert(0, self.pick(corpus.HEADING_EXTRAS[:3]))
        sec = 0
        for topic in topics:
            sec += 1
            b.append(("HEADING", self._number(st, sec, None) + topic))
            self._section_body(b, subj, subject_name)
            if self.r.random() < 0.6:
                for sub in range(1, self.r.randint(2, 3) + 1):
                    subtopic = self.pick(subj["topics"] + corpus.HEADING_EXTRAS)
                    b.append(("SUBHEADING", self._number(st, sec, sub) + subtopic))
                    self._section_body(b, subj, subject_name, short=True)
        if self.r.random() < 0.8:
            sec += 1
            b.append(("HEADING", self._number(st, sec, None) + self.pick(corpus.QUESTION_SECTION_TITLES)))
            for n in range(1, self.r.randint(2, 5) + 1):
                b.append(("QUESTION", self.question(subj, subject_name, n, st.question_prefix)))
                if self.r.random() < 0.6:
                    b.append(("ANSWER", self.answer(subj, subject_name, n)))
        if self.r.random() < 0.8:
            sec += 1
            b.append(("HEADING", self._number(st, sec, None) + self.pick(corpus.CONCLUSION_SECTION_TITLES)))
            b.append(("CONCLUSION", self.conclusion(subj, subject_name)))
        if self.r.random() < 0.75:
            b.append(("HEADING", self.pick(corpus.REFERENCE_SECTION_TITLES)))
            for i in range(self.r.randint(1, 3)):
                b.append(("REFERENCE", self.reference(subj, i, st.ref_style)))
        return spec

    def _number(self, st: Style, sec: int, sub: int | None) -> str:
        if st.heading_numbering == "none":
            return ""
        if sub is None:
            return f"Chapter {sec}: " if st.heading_numbering == "chapter" and self.r.random() < 0.3 else f"{sec}. "
        return f"{sec}.{sub} "

    def _section_body(self, b, subj, name, short: bool = False) -> None:
        b.append(("PARAGRAPH", self.paragraph(subj, name)))
        extras = ["DEFINITION", "EXAMPLE", "PROCEDURE", "IMPORTANT_POINT", "FORMULA", "PARAGRAPH"]
        for label in self.r.sample(extras, self.r.randint(1, 2 if short else 4)):
            if label == "DEFINITION":
                b.append((label, self.definition(subj, name)))
            elif label == "EXAMPLE":
                b.append((label, self.example(subj, name)))
            elif label == "PROCEDURE":
                b.append((label, self.procedure()))
            elif label == "IMPORTANT_POINT":
                b.append((label, self.important(subj, name)))
            elif label == "FORMULA" and subj["formulas"]:
                b.append((label, self.pick(subj["formulas"])))
            elif label == "PARAGRAPH":
                b.append((label, self.paragraph(subj, name)))


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------

_FONTS_REGISTERED = False


def _register_fonts() -> None:
    global _FONTS_REGISTERED
    if _FONTS_REGISTERED:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pairs = {
        "DejaVu": SYS_FONTS / "dejavu/DejaVuSans.ttf",
        "DejaVu-Bold": SYS_FONTS / "dejavu/DejaVuSans-Bold.ttf",
        "LibSans": SYS_FONTS / "liberation/LiberationSans-Regular.ttf",
        "LibSans-Bold": SYS_FONTS / "liberation/LiberationSans-Bold.ttf",
        "LibSerif": SYS_FONTS / "liberation/LiberationSerif-Regular.ttf",
        "LibSerif-Bold": SYS_FONTS / "liberation/LiberationSerif-Bold.ttf",
    }
    for name, path in pairs.items():
        if path.exists():
            pdfmetrics.registerFont(TTFont(name, str(path)))
    _FONTS_REGISTERED = True


def _font_names(family: str) -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics

    wanted = {"sans": ("LibSans", "LibSans-Bold"), "serif": ("LibSerif", "LibSerif-Bold"), "dejavu": ("DejaVu", "DejaVu-Bold")}[family]
    if wanted[0] in pdfmetrics.getRegisteredFontNames():
        return wanted
    return ("Helvetica", "Helvetica-Bold") if family != "serif" else ("Times-Roman", "Times-Bold")


def render_pdf(spec: DocSpec) -> bytes:
    from reportlab.lib.pagesizes import A4, LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    _register_fonts()
    st = spec.style
    regular, bold = _font_names(st.family)
    b = st.body_size

    def ps(name, font, size, before=0, after=6, indent=0):
        return ParagraphStyle(name, fontName=font, fontSize=size, leading=size * 1.3, spaceBefore=before, spaceAfter=after, leftIndent=indent)

    styles = {
        "TITLE": ps("t", bold, st.title_size, 0, 14),
        "HEADING": ps("h", bold if st.heading_bold else regular, st.heading_size, 10, 6),
        "SUBHEADING": ps("s", bold if st.sub_bold else regular, st.sub_size, 6, 4),
        "FORMULA": ps("f", regular, b, 2, 8, indent=30),
        "REFERENCE": ps("r", regular, max(b - 1.5, 8), 0, 3),
        "QUESTION": ps("q", regular, b, 2, 4),
    }
    body = ps("b", regular, b, 0, 8)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=random.Random(spec.doc_id).choice([A4, LETTER]))
    story = []
    for label, text in spec.blocks:
        t = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(t, styles.get(label, body)))
    doc.build(story)
    return buf.getvalue()


def render_docx(spec: DocSpec) -> bytes:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    for label, text in spec.blocks:
        if label == "TITLE":
            d.add_heading(text, level=0)
        elif label == "HEADING":
            d.add_heading(text, level=1)
        elif label == "SUBHEADING":
            d.add_heading(text, level=2)
        else:
            p = d.add_paragraph()
            run = p.add_run(text)
            run.font.size = Pt(spec.style.body_size)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_txt(spec: DocSpec) -> bytes:
    out = []
    for label, text in spec.blocks:
        if label == "TITLE" and random.Random(spec.doc_id).random() < 0.5:
            text = text.upper()
        out.append(text)
    return "\n\n".join(out).encode("utf-8")


# ---------------------------------------------------------------------------
# Alignment: extracted segments → source labels
# ---------------------------------------------------------------------------


def _norm(t: str) -> str:
    return re.sub(r"[^a-z0-9]", "", t.lower())


def align(spec: DocSpec, segments) -> list[tuple[object, str]]:
    """Sequentially match each extracted segment to the source block it came from."""
    src = [(lbl, _norm(t)) for lbl, t in spec.blocks]
    out = []
    j = 0
    for seg in segments:
        key = _norm(seg.text)
        if not key:
            continue
        for k in range(j, min(j + 4, len(src))):
            lbl, s = src[k]
            if key == s or (len(key) >= 12 and s.startswith(key[:40])) or (len(s) >= 12 and key.startswith(s[:40])):
                out.append((seg, lbl))
                j = k + 1
                break
    return out


def build_samples(spec: DocSpec) -> list[dict]:
    from backend.ml.features import segment_samples

    fmt = spec.style.fmt
    data = {"pdf": render_pdf, "docx": render_docx, "txt": render_txt}[fmt](spec)
    segments = preprocess(extract(fmt, data, allow_ocr=False))
    aligned = align(spec, segments)
    segs = [s for s, _ in aligned]
    samples = segment_samples(segs)
    for sample, (_, label) in zip(samples, aligned):
        sample["label"] = label
        sample["doc_id"] = spec.doc_id
        sample["subject"] = spec.subject
        sample["fmt"] = fmt
    return samples


# Subject-level split: every subject lives in exactly one split.
SPLITS = {
    "train": ["biology", "physics", "history", "economics", "computer_science", "psychology", "business", "environmental_science"],
    "val": ["chemistry", "geography"],
    "test": ["mathematics", "literature"],
}


def generate(docs_per_subject: int = 40, seed: int = 13) -> list[dict]:
    gen = Generator(seed)
    samples: list[dict] = []
    for subject in corpus.SUBJECTS:
        for i in range(docs_per_subject):
            fmt = gen.r.choices(["pdf", "docx", "txt"], weights=[0.7, 0.15, 0.15])[0]
            spec = gen.document(f"{subject}-{i:03d}", subject, fmt)
            samples.extend(build_samples(spec))
    return samples


def split_of(sample: dict) -> str:
    for name, subjects in SPLITS.items():
        if sample["subject"] in subjects:
            return name
    return "train"


def save_jsonl(samples: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        for s in samples:
            fh.write(json.dumps(s, ensure_ascii=False) + "\n")


def load_jsonl(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as fh:
        return [json.loads(line) for line in fh if line.strip()]
